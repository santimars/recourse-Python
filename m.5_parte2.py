import openpyxl
from openpyxl.chart import BarChart, Reference

# Crear un libro de trabajo
wb = openpyxl.Workbook()

# Seleccionar la hoja activa
sheet = wb.active
hojafirst = sheet.title
nombres = wb.sheetnames

print(nombres)

# Datos de ejemplo
data = [
    ['Categoria', 'Valor1', 'Valor2', 'Valor3'],
    ['A', 10, 15, 20],
    ['B', 5, 8, 12],
    ['C', 12, 6, 8],
    ['D', 23, 46, 18],
    ['E', 22, 16, 58]
]

# Agregar datos a la hoja de cálculo
for row in data:
    sheet.append(row)

# Crear un objeto de gráfico de barras
chart = BarChart()

# Configurar las series de datos y categorías
data_range = Reference(sheet, min_col=2, min_row=1, max_col=4, max_row=len(data))
categories = Reference(sheet, min_col=1, min_row=2, max_row=len(data))
chart.add_data(data_range, titles_from_data=True)
chart.set_categories(categories)

# Establecer el título del gráfico
chart.title = "Gráfico de Barras Avanzado"

# Establecer etiquetas de ejes
chart.x_axis.title = "Categorías"
chart.y_axis.title = "Valores"

# Agregar el gráfico a la hoja de cálculo
sheet.add_chart(chart, "H5")

# Guardar el libro de trabajo
wb.save("/Users/santirodriguez/Desktop/Libro1.xlsx")
############################################################################################
